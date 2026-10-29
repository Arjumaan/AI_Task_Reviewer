import os
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from playwright.sync_api import sync_playwright

output_dir = Path("templates")
output_dir.mkdir(parents=True, exist_ok=True)

md_content = """# CloudSync Engine — Core Microservices & Integration Task Specification

**Organization:** CloudSync Infrastructure Labs  
**Document Type:** Technical Assignment & Engineering Task Specification  
**Target Intern Role:** Backend & Cloud Engineering Intern  
**Target Deadline:** 7 days from initial Google Meet briefing session  
**Review Platform:** AI Task Reviewer Automated Evaluation Engine  

---

## 1. Executive Summary

During our Q4 reliability and security review of the CloudSync microservices architecture, several critical issues and feature enhancements were identified across four primary platform services: api.cloudsync.internal, auth.cloudsync.internal, worker.cloudsync.internal, and storage.cloudsync.internal. These microservices manage real-time file synchronization, user authentication, and high-throughput background processing for over 250,000 active users.

The purpose of this assignment is for the intern to investigate the root causes of these vulnerabilities and performance bottlenecks, implement clean, modular code fixes, and establish automated unit and integration tests. Upon completion, the intern will present their deliverables in a live review session where code quality, deterministic test coverage, and architecture decisions will be evaluated against our strict engineering rubrics.

---

## 2. Target Platforms & Scope

The scope of this technical task covers the following microservices:
- auth.cloudsync.internal: Authentication, session state, and OAuth2 security tokens.
- api.cloudsync.internal: Public REST API gateway and high-frequency endpoints.
- worker.cloudsync.internal: Celery and Redis background queue processing engine.
- storage.cloudsync.internal: S3-compatible chunked upload and file integrity subsystem.

---

## 3. Overall Findings & Requirements Snapshot

ID Platform Primary Area Severity Status
REQ-01 auth.cloudsync.internal JWT Refresh Token Revocation & Expiration Critical Open
REQ-02 api.cloudsync.internal Token Bucket Rate Limiting for Burst Traffic High Open
REQ-03 worker.cloudsync.internal Asynchronous Dead Letter Queue & Retry Backoff High Open
REQ-04 storage.cloudsync.internal SHA-256 Checksum Validation for Chunked Uploads Medium Open

---

## 4. Detailed Technical Specifications & Acceptance Criteria

### REQ-01 — JWT Refresh Token Revocation & Expiration
- Platform: auth.cloudsync.internal
- Severity: Critical
- Problem Statement: Refresh tokens currently lack explicit database revocation checks upon user logout or password changes, allowing stale tokens to generate valid access tokens indefinitely.
- Implementation Requirements:
  1. Store refresh token jti in Redis with active TTL.
  2. Implement an explicit /auth/revoke endpoint to blacklist revoked tokens immediately.
  3. Reject any refresh requests containing blacklisted or expired tokens with HTTP 401 Unauthorized.
- Acceptance Criteria (AC-01.1): 100% pass on deterministic integration tests verifying immediate token rejection after logout.

### REQ-02 — Token Bucket Rate Limiting for Burst Traffic
- Platform: api.cloudsync.internal
- Severity: High
- Problem Statement: High-volume client requests during sync surges can overwhelm downstream databases due to missing rate limiting headers and bucket controls.
- Implementation Requirements:
  1. Implement a distributed Token Bucket rate limiter middleware (limit: 100 requests per minute per client IP).
  2. Return standard HTTP headers (X-RateLimit-Limit, X-RateLimit-Remaining, X-RateLimit-Reset).
  3. Gracefully return HTTP 429 Too Many Requests when limits are exceeded.
- Acceptance Criteria (AC-02.1): Deterministic mock tests simulate 150 concurrent requests and verify 50 requests receive HTTP 429 with correct headers.

### REQ-03 — Asynchronous Dead Letter Queue & Retry Backoff
- Platform: worker.cloudsync.internal
- Severity: High
- Problem Statement: Transient network errors during external webhooks cause worker task drops without retry backoff or quarantine.
- Implementation Requirements:
  1. Implement exponential backoff retry mechanism (maximum 3 retries, base delay: 2s).
  2. Route persistently failed tasks to a dedicated Dead Letter Queue (DLQ) with error telemetry.
  3. Emit structured logs containing error details and original payload metadata.
- Acceptance Criteria (AC-03.1): Automated worker tests verify 3 retry attempts before routing failed jobs to DLQ.

### REQ-04 — SHA-256 Checksum Validation for Chunked Uploads
- Platform: storage.cloudsync.internal
- Severity: Medium
- Problem Statement: Multi-part uploaded file chunks are stitched without final hash verification, risking silent data corruption during packet loss.
- Implementation Requirements:
  1. Calculate rolling SHA-256 checksum during chunk reassembly.
  2. Compare calculated hash against client-provided Content-SHA256 header.
  3. Abort assembly and purge corrupted chunk artifacts if hash mismatch occurs.
- Acceptance Criteria (AC-04.1): Artifact inspection test verifying corrupt chunk payload triggers immediate cleanup and HTTP 422 Unprocessable Entity.

---

## 5. Expected Intern Deliverables

To pass the technical evaluation, the intern must deliver:
1. Source Code Patches: Modular, well-formatted code addressing REQ-01 through REQ-04 with zero regressions.
2. Automated Test Suite: Unit and integration test suite with minimum 90% code coverage.
3. Pull Request & Repository: Submitted GitHub Pull Request link with commit hash and description.
4. Verification Report: Reproduction steps, test execution logs, and sandbox instructions.

---

## 6. Evaluation Criteria & Review Protocol

- Deterministic Testing (Weight: 40%): Automated execution of all test suites against staging criteria.
- Code Quality & Architecture (Weight: 30%): Clean separation of concerns, error handling, and type safety.
- Security & Integrity (Weight: 20%): Adherence to authentication standards and defense-in-depth principles.
- Documentation & Communication (Weight: 10%): Clear PR notes and verbal defense during the live Google Meet review.
"""

# 1. Write Markdown & TXT
md_path = output_dir / "SAMPLE_TASK_SPECIFICATION_TEMPLATE.md"
with open(md_path, "w", encoding="utf-8") as f:
    f.write(md_content)
print(f"Saved: {md_path}")

txt_path = output_dir / "SAMPLE_TASK_SPECIFICATION_TEMPLATE.txt"
with open(txt_path, "w", encoding="utf-8") as f:
    f.write(md_content)
print(f"Saved: {txt_path}")

# 2. Write DOCX
doc = Document()
doc.add_heading("CloudSync Engine — Core Microservices & Integration Task Specification", 0)

p_meta = doc.add_paragraph()
p_meta.add_run("Organization: ").bold = True
p_meta.add_run("CloudSync Infrastructure Labs\n")
p_meta.add_run("Document Type: ").bold = True
p_meta.add_run("Technical Assignment & Engineering Task Specification\n")
p_meta.add_run("Target Intern Role: ").bold = True
p_meta.add_run("Backend & Cloud Engineering Intern\n")
p_meta.add_run("Target Deadline: ").bold = True
p_meta.add_run("7 days from initial Google Meet briefing session\n")
p_meta.add_run("Review Platform: ").bold = True
p_meta.add_run("AI Task Reviewer Automated Evaluation Engine")

doc.add_heading("1. Executive Summary", level=1)
doc.add_paragraph(
    "During our Q4 reliability and security review of the CloudSync microservices architecture, several critical issues and feature enhancements were identified across four primary platform services: api.cloudsync.internal, auth.cloudsync.internal, worker.cloudsync.internal, and storage.cloudsync.internal. These microservices manage real-time file synchronization, user authentication, and high-throughput background processing for over 250,000 active users."
)
doc.add_paragraph(
    "The purpose of this assignment is for the intern to investigate the root causes of these vulnerabilities and performance bottlenecks, implement clean, modular code fixes, and establish automated unit and integration tests. Upon completion, the intern will present their deliverables in a live review session where code quality, deterministic test coverage, and architecture decisions will be evaluated against our strict engineering rubrics."
)

doc.add_heading("2. Target Platforms & Scope", level=1)
doc.add_paragraph("The scope of this technical task covers the following microservices:", style='List Bullet')
doc.add_paragraph("auth.cloudsync.internal: Authentication, session state, and OAuth2 security tokens.", style='List Bullet')
doc.add_paragraph("api.cloudsync.internal: Public REST API gateway and high-frequency endpoints.", style='List Bullet')
doc.add_paragraph("worker.cloudsync.internal: Celery and Redis background queue processing engine.", style='List Bullet')
doc.add_paragraph("storage.cloudsync.internal: S3-compatible chunked upload and file integrity subsystem.", style='List Bullet')

doc.add_heading("3. Overall Findings & Requirements Snapshot", level=1)
table = doc.add_table(rows=1, cols=5)
hdr_cells = table.rows[0].cells
headers = ["ID", "Platform", "Primary Area", "Severity", "Status"]
for i, h in enumerate(headers):
    hdr_cells[i].text = h
    hdr_cells[i].paragraphs[0].runs[0].bold = True

req_rows = [
    ("REQ-01", "auth.cloudsync.internal", "JWT Refresh Token Revocation & Expiration", "Critical", "Open"),
    ("REQ-02", "api.cloudsync.internal", "Token Bucket Rate Limiting for Burst Traffic", "High", "Open"),
    ("REQ-03", "worker.cloudsync.internal", "Asynchronous Dead Letter Queue & Retry Backoff", "High", "Open"),
    ("REQ-04", "storage.cloudsync.internal", "SHA-256 Checksum Validation for Chunked Uploads", "Medium", "Open")
]
for row in req_rows:
    row_cells = table.add_row().cells
    for i, val in enumerate(row):
        row_cells[i].text = val

doc.add_heading("4. Detailed Technical Specifications & Acceptance Criteria", level=1)
specs = [
    ("REQ-01 — JWT Refresh Token Revocation & Expiration", "auth.cloudsync.internal", "Critical",
     "Refresh tokens currently lack explicit database revocation checks upon user logout or password changes, allowing stale tokens to generate valid access tokens indefinitely.",
     ["Store refresh token jti in Redis with active TTL.", "Implement an explicit /auth/revoke endpoint to blacklist revoked tokens immediately.", "Reject any refresh requests containing blacklisted or expired tokens with HTTP 401 Unauthorized."],
     "100% pass on deterministic integration tests verifying immediate token rejection after logout."),
    ("REQ-02 — Token Bucket Rate Limiting for Burst Traffic", "api.cloudsync.internal", "High",
     "High-volume client requests during sync surges can overwhelm downstream databases due to missing rate limiting headers and bucket controls.",
     ["Implement a distributed Token Bucket rate limiter middleware (limit: 100 requests per minute per client IP).", "Return standard HTTP headers (X-RateLimit-Limit, X-RateLimit-Remaining, X-RateLimit-Reset).", "Gracefully return HTTP 429 Too Many Requests when limits are exceeded."],
     "Deterministic mock tests simulate 150 concurrent requests and verify 50 requests receive HTTP 429 with correct headers."),
    ("REQ-03 — Asynchronous Dead Letter Queue & Retry Backoff", "worker.cloudsync.internal", "High",
     "Transient network errors during external webhooks cause worker task drops without retry backoff or quarantine.",
     ["Implement exponential backoff retry mechanism (maximum 3 retries, base delay: 2s).", "Route persistently failed tasks to a dedicated Dead Letter Queue (DLQ) with error telemetry.", "Emit structured logs containing error details and original payload metadata."],
     "Automated worker tests verify 3 retry attempts before routing failed jobs to DLQ."),
    ("REQ-04 — SHA-256 Checksum Validation for Chunked Uploads", "storage.cloudsync.internal", "Medium",
     "Multi-part uploaded file chunks are stitched without final hash verification, risking silent data corruption during packet loss.",
     ["Calculate rolling SHA-256 checksum during chunk reassembly.", "Compare calculated hash against client-provided Content-SHA256 header.", "Abort assembly and purge corrupted chunk artifacts if hash mismatch occurs."],
     "Artifact inspection test verifying corrupt chunk payload triggers immediate cleanup and HTTP 422 Unprocessable Entity.")
]

for title, plat, sev, prob, reqs, ac in specs:
    doc.add_heading(title, level=2)
    p = doc.add_paragraph()
    p.add_run(f"Platform: {plat} | Severity: {sev}\n").bold = True
    p.add_run(f"Problem Statement: {prob}\n")
    doc.add_paragraph("Implementation Requirements:")
    for r in reqs:
        doc.add_paragraph(r, style='List Bullet')
    p_ac = doc.add_paragraph()
    p_ac.add_run(f"Acceptance Criteria: {ac}").italic = True

doc.add_heading("5. Expected Intern Deliverables", level=1)
deliverables = [
    "Source Code Patches: Modular, well-formatted code addressing REQ-01 through REQ-04 with zero regressions.",
    "Automated Test Suite: Unit and integration test suite with minimum 90% code coverage.",
    "Pull Request & Repository: Submitted GitHub Pull Request link with commit hash and description.",
    "Verification Report: Reproduction steps, test execution logs, and sandbox instructions."
]
for d in deliverables:
    doc.add_paragraph(d, style='List Bullet')

doc.add_heading("6. Evaluation Criteria & Review Protocol", level=1)
rubrics = [
    "Deterministic Testing (Weight: 40%): Automated execution of all test suites against staging criteria.",
    "Code Quality & Architecture (Weight: 30%): Clean separation of concerns, error handling, and type safety.",
    "Security & Integrity (Weight: 20%): Adherence to authentication standards and defense-in-depth principles.",
    "Documentation & Communication (Weight: 10%): Clear PR notes and verbal defense during the live Google Meet review."
]
for rb in rubrics:
    doc.add_paragraph(rb, style='List Bullet')

docx_path = output_dir / "SAMPLE_TASK_SPECIFICATION_TEMPLATE.docx"
doc.save(str(docx_path))
print(f"Saved: {docx_path}")

# 3. Write Styled PDF using Playwright
html_content = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
  body {{
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
    color: #1e293b;
    line-height: 1.5;
    padding: 30px;
    font-size: 11pt;
  }}
  h1 {{
    color: #0f172a;
    font-size: 18pt;
    border-bottom: 2px solid #0284c7;
    padding-bottom: 6px;
    margin-bottom: 12px;
  }}
  h2 {{
    color: #0369a1;
    font-size: 13pt;
    margin-top: 18px;
    margin-bottom: 8px;
    border-bottom: 1px solid #e2e8f0;
    padding-bottom: 4px;
  }}
  h3 {{
    color: #0f172a;
    font-size: 11pt;
    margin-top: 12px;
    margin-bottom: 4px;
  }}
  .meta-box {{
    background: #f8fafc;
    border: 1px solid #cbd5e1;
    border-radius: 6px;
    padding: 10px 14px;
    margin-bottom: 16px;
    font-size: 10pt;
  }}
  .meta-row {{
    display: flex;
    margin-bottom: 4px;
  }}
  .meta-label {{
    font-weight: bold;
    width: 140px;
    color: #475569;
  }}
  table {{
    width: 100%;
    border-collapse: collapse;
    margin: 12px 0;
    font-size: 10pt;
  }}
  th, td {{
    border: 1px solid #cbd5e1;
    padding: 6px 10px;
    text-align: left;
  }}
  th {{
    background: #f1f5f9;
    font-weight: 600;
    color: #334155;
  }}
  .badge-crit {{
    background: #ffe4e6;
    color: #9f1239;
    padding: 2px 6px;
    border-radius: 4px;
    font-weight: bold;
    font-size: 9pt;
  }}
  .badge-high {{
    background: #fef3c7;
    color: #92400e;
    padding: 2px 6px;
    border-radius: 4px;
    font-weight: bold;
    font-size: 9pt;
  }}
  .badge-med {{
    background: #e0f2fe;
    color: #0369a1;
    padding: 2px 6px;
    border-radius: 4px;
    font-weight: bold;
    font-size: 9pt;
  }}
  ul {{
    margin-top: 4px;
    margin-bottom: 8px;
    padding-left: 20px;
  }}
  li {{
    margin-bottom: 3px;
  }}
  .ac-box {{
    background: #f0fdf4;
    border-left: 3px solid #16a34a;
    padding: 4px 8px;
    margin: 4px 0 10px 0;
    font-size: 9.5pt;
    color: #14532d;
  }}
</style>
</head>
<body>
  <h1>CloudSync Engine — Core Microservices & Integration Task Specification</h1>
  
  <div class="meta-box">
    <div class="meta-row"><span class="meta-label">Organization:</span><span>CloudSync Infrastructure Labs</span></div>
    <div class="meta-row"><span class="meta-label">Document Type:</span><span>Technical Assignment & Engineering Task Specification</span></div>
    <div class="meta-row"><span class="meta-label">Target Intern Role:</span><span>Backend & Cloud Engineering Intern</span></div>
    <div class="meta-row"><span class="meta-label">Target Deadline:</span><span>7 days from initial Google Meet briefing session</span></div>
    <div class="meta-row"><span class="meta-label">Review Platform:</span><span>AI Task Reviewer Automated Evaluation Engine</span></div>
  </div>

  <h2>1. Executive Summary</h2>
  <p>During our Q4 reliability and security review of the CloudSync microservices architecture, several critical issues and feature enhancements were identified across four primary platform services: <code>api.cloudsync.internal</code>, <code>auth.cloudsync.internal</code>, <code>worker.cloudsync.internal</code>, and <code>storage.cloudsync.internal</code>. These microservices manage real-time file synchronization, user authentication, and high-throughput background processing for over 250,000 active users.</p>
  <p>The purpose of this assignment is for the intern to investigate the root causes of these vulnerabilities and performance bottlenecks, implement clean, modular code fixes, and establish automated unit and integration tests. Upon completion, the intern will present their deliverables in a live review session where code quality, deterministic test coverage, and architecture decisions will be evaluated against our strict engineering rubrics.</p>

  <h2>2. Target Platforms & Scope</h2>
  <ul>
    <li><strong>auth.cloudsync.internal</strong>: Authentication, session state, and OAuth2 security tokens.</li>
    <li><strong>api.cloudsync.internal</strong>: Public REST API gateway and high-frequency endpoints.</li>
    <li><strong>worker.cloudsync.internal</strong>: Celery and Redis background queue processing engine.</li>
    <li><strong>storage.cloudsync.internal</strong>: S3-compatible chunked upload and file integrity subsystem.</li>
  </ul>

  <h2>3. Overall Findings & Requirements Snapshot</h2>
  <table>
    <thead>
      <tr>
        <th>ID</th>
        <th>Platform</th>
        <th>Primary Area</th>
        <th>Severity</th>
        <th>Status</th>
      </tr>
    </thead>
    <tbody>
      <tr>
        <td><strong>REQ-01</strong></td>
        <td>auth.cloudsync.internal</td>
        <td>JWT Refresh Token Revocation & Expiration</td>
        <td><span class="badge-crit">Critical</span></td>
        <td>Open</td>
      </tr>
      <tr>
        <td><strong>REQ-02</strong></td>
        <td>api.cloudsync.internal</td>
        <td>Token Bucket Rate Limiting for Burst Traffic</td>
        <td><span class="badge-high">High</span></td>
        <td>Open</td>
      </tr>
      <tr>
        <td><strong>REQ-03</strong></td>
        <td>worker.cloudsync.internal</td>
        <td>Asynchronous Dead Letter Queue & Retry Backoff</td>
        <td><span class="badge-high">High</span></td>
        <td>Open</td>
      </tr>
      <tr>
        <td><strong>REQ-04</strong></td>
        <td>storage.cloudsync.internal</td>
        <td>SHA-256 Checksum Validation for Chunked Uploads</td>
        <td><span class="badge-med">Medium</span></td>
        <td>Open</td>
      </tr>
    </tbody>
  </table>

  <h2>4. Detailed Technical Specifications & Acceptance Criteria</h2>
  
  <h3>REQ-01 — JWT Refresh Token Revocation & Expiration</h3>
  <p><strong>Platform:</strong> auth.cloudsync.internal | <strong>Severity:</strong> <span class="badge-crit">Critical</span><br>
  <strong>Problem Statement:</strong> Refresh tokens currently lack explicit database revocation checks upon user logout or password changes, allowing stale tokens to generate valid access tokens indefinitely.</p>
  <ul>
    <li>Store refresh token jti in Redis with active TTL.</li>
    <li>Implement an explicit <code>/auth/revoke</code> endpoint to blacklist revoked tokens immediately.</li>
    <li>Reject any refresh requests containing blacklisted or expired tokens with HTTP 401 Unauthorized.</li>
  </ul>
  <div class="ac-box"><strong>Acceptance Criteria (AC-01.1):</strong> 100% pass on deterministic integration tests verifying immediate token rejection after logout.</div>

  <h3>REQ-02 — Token Bucket Rate Limiting for Burst Traffic</h3>
  <p><strong>Platform:</strong> api.cloudsync.internal | <strong>Severity:</strong> <span class="badge-high">High</span><br>
  <strong>Problem Statement:</strong> High-volume client requests during sync surges can overwhelm downstream databases due to missing rate limiting headers and bucket controls.</p>
  <ul>
    <li>Implement a distributed Token Bucket rate limiter middleware (limit: 100 requests per minute per client IP).</li>
    <li>Return standard HTTP headers (<code>X-RateLimit-Limit</code>, <code>X-RateLimit-Remaining</code>, <code>X-RateLimit-Reset</code>).</li>
    <li>Gracefully return HTTP 429 Too Many Requests when limits are exceeded.</li>
  </ul>
  <div class="ac-box"><strong>Acceptance Criteria (AC-02.1):</strong> Deterministic mock tests simulate 150 concurrent requests and verify 50 requests receive HTTP 429 with correct headers.</div>

  <h3>REQ-03 — Asynchronous Dead Letter Queue & Retry Backoff</h3>
  <p><strong>Platform:</strong> worker.cloudsync.internal | <strong>Severity:</strong> <span class="badge-high">High</span><br>
  <strong>Problem Statement:</strong> Transient network errors during external webhooks cause worker task drops without retry backoff or quarantine.</p>
  <ul>
    <li>Implement exponential backoff retry mechanism (maximum 3 retries, base delay: 2s).</li>
    <li>Route persistently failed tasks to a dedicated Dead Letter Queue (DLQ) with error telemetry.</li>
    <li>Emit structured logs containing error details and original payload metadata.</li>
  </ul>
  <div class="ac-box"><strong>Acceptance Criteria (AC-03.1):</strong> Automated worker tests verify 3 retry attempts before routing failed jobs to DLQ.</div>

  <h3>REQ-04 — SHA-256 Checksum Validation for Chunked Uploads</h3>
  <p><strong>Platform:</strong> storage.cloudsync.internal | <strong>Severity:</strong> <span class="badge-med">Medium</span><br>
  <strong>Problem Statement:</strong> Multi-part uploaded file chunks are stitched without final hash verification, risking silent data corruption during packet loss.</p>
  <ul>
    <li>Calculate rolling SHA-256 checksum during chunk reassembly.</li>
    <li>Compare calculated hash against client-provided <code>Content-SHA256</code> header.</li>
    <li>Abort assembly and purge corrupted chunk artifacts if hash mismatch occurs.</li>
  </ul>
  <div class="ac-box"><strong>Acceptance Criteria (AC-04.1):</strong> Artifact inspection test verifying corrupt chunk payload triggers immediate cleanup and HTTP 422 Unprocessable Entity.</div>

  <h2>5. Expected Intern Deliverables</h2>
  <ul>
    <li><strong>Source Code Patches:</strong> Modular, well-formatted code addressing REQ-01 through REQ-04 with zero regressions.</li>
    <li><strong>Automated Test Suite:</strong> Unit and integration test suite with minimum 90% code coverage.</li>
    <li><strong>Pull Request & Repository:</strong> Submitted GitHub Pull Request link with commit hash and description.</li>
    <li><strong>Verification Report:</strong> Reproduction steps, test execution logs, and sandbox instructions.</li>
  </ul>

  <h2>6. Evaluation Criteria & Review Protocol</h2>
  <ul>
    <li><strong>Deterministic Testing (Weight: 40%):</strong> Automated execution of all test suites against staging criteria.</li>
    <li><strong>Code Quality & Architecture (Weight: 30%):</strong> Clean separation of concerns, error handling, and type safety.</li>
    <li><strong>Security & Integrity (Weight: 20%):</strong> Adherence to authentication standards and defense-in-depth principles.</li>
    <li><strong>Documentation & Communication (Weight: 10%):</strong> Clear PR notes and verbal defense during the live Google Meet review.</li>
  </ul>
</body>
</html>
"""

pdf_path = output_dir / "SAMPLE_TASK_SPECIFICATION_TEMPLATE.pdf"
with sync_playwright() as p:
    browser = p.chromium.launch()
    page = browser.new_page()
    page.set_content(html_content)
    page.pdf(path=str(pdf_path), format="A4", print_background=True, margin={"top": "20mm", "bottom": "20mm", "left": "15mm", "right": "15mm"})
    browser.close()

print(f"Saved: {pdf_path}")
print("All 4 template formats generated successfully.")
