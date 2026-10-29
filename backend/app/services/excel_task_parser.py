"""
Excel/CSV Task Assignment Parser.
Parses spreadsheets from HR with per-intern task assignments.
"""
import csv
import io
import logging
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional

logger = logging.getLogger("task_reviewer.excel_parser")

class ExcelTaskParser:

    REQUIRED_COLUMNS = [
        "intern email", "intern name", "task title",
        "task objective / requirements", "deadline date", "meeting time"
    ]

    @classmethod
    def parse_xlsx(cls, file_path: str) -> List[Dict[str, Any]]:
        """Parse an Excel (.xlsx) file into task assignment rows."""
        import openpyxl
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        ws = wb.active
        if ws is None:
            raise ValueError("Excel file has no active worksheet.")

        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            raise ValueError("Excel file must have a header row and at least one data row.")

        headers = [str(h).strip().lower() if h else "" for h in rows[0]]
        missing = [col for col in cls.REQUIRED_COLUMNS if col not in headers]
        if missing:
            raise ValueError(f"Missing required columns: {', '.join(missing)}. Expected: {', '.join(cls.REQUIRED_COLUMNS)}")

        col_map = {col: headers.index(col) for col in cls.REQUIRED_COLUMNS}
        results = []
        for row_idx, row in enumerate(rows[1:], start=2):
            email = str(row[col_map["intern email"]] or "").strip()
            name = str(row[col_map["intern name"]] or "").strip()
            title = str(row[col_map["task title"]] or "").strip()
            objective = str(row[col_map["task objective / requirements"]] or "").strip()
            deadline_raw = row[col_map["deadline date"]]
            meeting_raw = row[col_map["meeting time"]]

            if not email or not title:
                logger.warning(f"[EXCEL PARSER] Row {row_idx} skipped: missing email or title.")
                continue

            if "@" not in email:
                logger.warning(f"[EXCEL PARSER] Row {row_idx} skipped: invalid email '{email}'.")
                continue

            deadline_str = cls._parse_date(deadline_raw)
            meeting_str = cls._parse_date(meeting_raw)

            results.append({
                "email": email.lower(),
                "name": name or email.split("@")[0].title(),
                "title": title,
                "objective": objective,
                "deadline": deadline_str,
                "meeting_time": meeting_str,
                "row_number": row_idx
            })

        wb.close()
        return results

    @classmethod
    def parse_csv(cls, file_content: bytes) -> List[Dict[str, Any]]:
        """Parse a CSV file (bytes) into task assignment rows."""
        text = file_content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))

        if not reader.fieldnames:
            raise ValueError("CSV file has no header row.")

        headers_lower = [h.strip().lower() for h in reader.fieldnames]
        missing = [col for col in cls.REQUIRED_COLUMNS if col not in headers_lower]
        if missing:
            raise ValueError(f"Missing required columns: {', '.join(missing)}.")

        field_map = {}
        for original, lower in zip(reader.fieldnames, headers_lower):
            for col in cls.REQUIRED_COLUMNS:
                if lower == col:
                    field_map[col] = original

        results = []
        for row_idx, row in enumerate(reader, start=2):
            email = (row.get(field_map.get("intern email", ""), "") or "").strip()
            name = (row.get(field_map.get("intern name", ""), "") or "").strip()
            title = (row.get(field_map.get("task title", ""), "") or "").strip()
            objective = (row.get(field_map.get("task objective / requirements", ""), "") or "").strip()
            deadline_raw = (row.get(field_map.get("deadline date", ""), "") or "").strip()
            meeting_raw = (row.get(field_map.get("meeting time", ""), "") or "").strip()

            if not email or not title or "@" not in email:
                continue

            results.append({
                "email": email.lower(),
                "name": name or email.split("@")[0].title(),
                "title": title,
                "objective": objective,
                "deadline": deadline_raw,
                "meeting_time": meeting_raw,
                "row_number": row_idx
            })

        return results

    @classmethod
    def _parse_date(cls, value: Any) -> Optional[str]:
        """Convert various date formats to ISO string."""
        if value is None:
            return None
        if isinstance(value, datetime):
            return value.isoformat()
        s = str(value).strip()
        if not s:
            return None
        for fmt in ["%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y", "%d/%m/%Y", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M"]:
            try:
                return datetime.strptime(s, fmt).isoformat()
            except ValueError:
                continue
        return s  # Return raw string if no format matched

    @classmethod
    def generate_template(cls, output_path: str) -> str:
        """Generate a downloadable Excel template with correct headers and sample data."""
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Task Assignments"

        headers = ["Intern Email", "Intern Name", "Task Title", "Task Objective / Requirements", "Deadline Date", "Meeting Time"]
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", bold=True, color="38BDF8", size=11)

        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        sample_rows = [
            ["intern1@company.com", "Aarav Sharma", "REST API Development", "Build a RESTful API with authentication, CRUD operations, and proper error handling using FastAPI or Express.js", "2026-11-15", "2026-11-02 10:00"],
            ["intern2@company.com", "Priya Menon", "React Dashboard UI", "Create a responsive admin dashboard with charts, data tables, and dark mode using React and Tailwind CSS", "2026-11-20", "2026-11-03 14:00"],
        ]
        for row_idx, row_data in enumerate(sample_rows, 2):
            for col_idx, value in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)

        col_widths = [28, 20, 30, 70, 16, 20]
        for col_idx, width in enumerate(col_widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

        wb.save(output_path)
        return output_path
